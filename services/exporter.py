"""导出清洗后的数据为 Excel。"""
from __future__ import annotations

import io
import re
from datetime import datetime
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="E8ECEF")
HEADER_FONT = Font(name="Microsoft YaHei", bold=True, size=10, color="1F2933")
BODY_FONT = Font(name="Microsoft YaHei", size=10, color="1F2933")
BORDER = Border(
    left=Side(style="thin", color="D5DAE0"),
    right=Side(style="thin", color="D5DAE0"),
    top=Side(style="thin", color="D5DAE0"),
    bottom=Side(style="thin", color="D5DAE0"),
)


def _safe_name(name: str) -> str:
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", name or "")
    name = re.sub(r"\s+", " ", name).strip()
    return (name or "数据集")[:60]


def _display_width(value: Any) -> int:
    text = "" if value is None else str(value)
    width = 0
    for char in text:
        width += 2 if ord(char) > 127 else 1
    return width


def _style_data_sheet(worksheet, frame: pd.DataFrame) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.row_dimensions[1].height = 24
    if worksheet.max_row >= 1:
        worksheet.auto_filter.ref = worksheet.dimensions

    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

    for index, column in enumerate(frame.columns, start=1):
        letter = get_column_letter(index)
        header_text = str(column)
        max_width = _display_width(header_text) + 4
        sample = frame[column].head(200)
        for value in sample:
            if value is None or (isinstance(value, float) and pd.isna(value)):
                continue
            max_width = max(max_width, _display_width(value) + 3)
        worksheet.column_dimensions[letter].width = max(8, min(42, max_width))


def _style_info_sheet(worksheet, rows: int) -> None:
    worksheet.column_dimensions["A"].width = 18
    worksheet.column_dimensions["B"].width = 86
    worksheet.sheet_view.showGridLines = False
    for row in worksheet.iter_rows(min_row=1, max_row=rows, min_col=1, max_col=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in worksheet[1]:
        cell.font = Font(name="Microsoft YaHei", bold=True, size=14, color="1F2933")
    worksheet.row_dimensions[1].height = 28


def build_excel(
    frame: pd.DataFrame,
    meta: dict[str, Any],
    *,
    scope: str,
    filters_text: str,
    sorts_text: str,
) -> tuple[io.BytesIO, str]:
    scope_label = "当前筛选视图" if scope == "view" else "全部清洗数据"
    buffer = io.BytesIO()

    with pd.ExcelWriter(
        buffer,
        engine="openpyxl",
        datetime_format="yyyy-mm-dd hh:mm:ss",
        date_format="yyyy-mm-dd",
    ) as writer:
        frame.to_excel(writer, sheet_name="数据", index=False)
        workbook = writer.book
        workbook.properties.creator = "数据洞察 - 轻量级数据可视化分析工具"
        workbook.properties.title = f"{meta.get('name', '数据集')} 数据导出"
        _style_data_sheet(writer.sheets["数据"], frame)

        info_rows = [
            (f"{meta.get('name', '数据集')} · 数据导出", ""),
            ("源文件", meta.get("file_name", "")),
            ("工作表", meta.get("sheet") or "—"),
            ("数据版本", f"v{meta.get('version', 0)}"),
            ("导出时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            ("导出范围", scope_label),
            ("筛选条件", filters_text or "无"),
            ("排序方式", sorts_text or "无"),
            ("行数", len(frame)),
            ("列数", len(frame.columns)),
            ("缺失单元格", int(frame.isna().sum().sum())),
            ("说明", "本文件由「数据洞察」导出，仅包含数据工作表与说明信息。"),
        ]
        sheet = workbook.create_sheet("导出说明")
        for row_index, (label, value) in enumerate(info_rows, start=1):
            sheet.cell(row=row_index, column=1, value=label)
            sheet.cell(row=row_index, column=2, value=value)
        _style_info_sheet(sheet, len(info_rows))
        workbook.active = 0

    buffer.seek(0)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M")
    suffix = "当前视图" if scope == "view" else "全部数据"
    filename = f"{_safe_name(str(meta.get('name', '数据集')))}_导出{suffix}_{timestamp}.xlsx"
    return buffer, filename


def build_chart_filename(meta: dict[str, Any], chart_type_label: str) -> str:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"{_safe_name(str(meta.get('name', '数据集')))}_{chart_type_label}_{timestamp}.png"
