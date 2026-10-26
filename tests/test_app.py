"""端到端测试：导入、预览、清洗、筛选、图表、导出、历史记录。"""
from __future__ import annotations

import io
import os
import shutil
import tempfile
import unittest
from pathlib import Path

TEST_DATA_DIR = tempfile.mkdtemp(prefix="dv_test_")
os.environ["DV_DATA_DIR"] = TEST_DATA_DIR
os.environ["DV_SECRET_KEY"] = "test-secret"

import app as application  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from PIL import Image  # noqa: E402

CSV_TEXT = """日期,地区,产品,销量,单价,销售额,编号,备注
2024-01-01,华东,键盘,10,99.5,995,001,首单
2024-01-02,华北,鼠标,20,49,980,002,
2024-01-03,华东,显示器,5,899,4495,003,大客户
2024-01-02,华北,鼠标,20,49,980,002,
2024/1/4,华南,键盘,,99.5,,004,
2024-01-05,华南,显示器,3,899,2697,005,加急
"""


def csv_bytes(encoding: str = "gbk") -> bytes:
    return CSV_TEXT.encode(encoding)


def upload_file(client, content: bytes, filename: str):
    return client.post(
        "/api/files/inspect",
        data={"file": (io.BytesIO(content), filename)},
        content_type="multipart/form-data",
    )


class DataVizTestCase(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        application.app.config.update(TESTING=True)
        cls.client = application.app.test_client()

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(TEST_DATA_DIR, ignore_errors=True)

    # ------------------------------------------------------------------
    def import_csv(self, name="销售测试"):
        response = upload_file(self.client, csv_bytes(), "销售测试.csv")
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        file_info = response.get_json()["file"]
        self.assertTrue(file_info["encoding"].lower().startswith("gb"))
        self.assertEqual(file_info["kind"], "csv")

        response = self.client.post(
            "/api/files/import",
            json={
                "token": file_info["token"],
                "file_name": "销售测试.csv",
                "name": name,
                "has_header": True,
                "header_row": 1,
                "encoding": file_info["encoding"],
                "delimiter": ",",
            },
        )
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        payload = response.get_json()
        self.assertTrue(payload["ok"])
        return payload["dataset_id"], file_info

    def summary(self, dataset_id):
        response = self.client.get(f"/api/datasets/{dataset_id}/summary")
        self.assertEqual(response.status_code, 200)
        return response.get_json()["summary"]

    # ------------------------------------------------------------------
    def test_01_import_and_preview(self):
        dataset_id, file_info = self.import_csv()
        summary = self.summary(dataset_id)
        self.assertEqual(summary["base_rows"], 6)
        self.assertEqual(summary["view_rows"], 6)
        self.assertEqual(summary["meta"]["column_count"], 8)

        columns = {item["name"]: item for item in summary["columns"]}
        self.assertEqual(columns["日期"]["type"], "日期时间")
        self.assertEqual(columns["销量"]["type"], "整数")
        self.assertEqual(columns["编号"]["type"], "文本")
        self.assertEqual(columns["备注"]["type"], "文本")

        profiles = {item["name"]: item for item in summary["profiles"]}
        self.assertEqual(profiles["销量"]["missing"], 1)
        self.assertEqual(profiles["备注"]["missing"], 3)
        self.assertAlmostEqual(profiles["销量"]["mean"], 11.6, places=4)
        self.assertEqual(profiles["销量"]["min"], 3)
        self.assertEqual(profiles["销量"]["max"], 20)

        response = self.client.get(f"/api/datasets/{dataset_id}/rows?page=1&page_size=10")
        rows = response.get_json()
        self.assertEqual(rows["total"], 6)
        self.assertEqual(rows["rows"][0]["编号"], "001")
        self.assertIsNone(rows["rows"][4]["销量"])

        response = self.client.get(f"/workspace/{dataset_id}")
        self.assertEqual(response.status_code, 200)
        html = response.get_data(as_text=True)
        self.assertIn("数据预览", html)
        self.assertIn("字段统计", html)

    def test_02_excel_sheet_selection(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "明细"
        sheet.append(["姓名", "分数"])
        sheet.append(["张三", 90])
        sheet.append(["李四", 85])
        extra = workbook.create_sheet("汇总")
        extra.append(["班级", "人数"])
        extra.append(["一班", 40])
        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)

        response = upload_file(
            self.client, buffer.getvalue(), "成绩表.xlsx"
        )
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        info = response.get_json()["file"]
        self.assertEqual(info["sheets"], ["明细", "汇总"])

        response = self.client.post(
            "/api/files/import",
            json={
                "token": info["token"],
                "file_name": "成绩表.xlsx",
                "name": "成绩表",
                "sheet": "汇总",
                "has_header": True,
                "header_row": 1,
            },
        )
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        dataset_id = response.get_json()["dataset_id"]
        summary = self.summary(dataset_id)
        self.assertEqual(summary["base_rows"], 1)
        self.assertEqual([item["name"] for item in summary["columns"]], ["班级", "人数"])
        self.assertEqual(summary["meta"]["sheet"], "汇总")

    def test_03_cleaning_fill_drop_convert_undo(self):
        dataset_id, _ = self.import_csv("清洗测试")

        response = self.client.post(
            f"/api/datasets/{dataset_id}/clean", json={"action": "dedup", "params": {}}
        )
        self.assertTrue(response.get_json()["ok"])
        self.assertEqual(self.summary(dataset_id)["base_rows"], 5)

        response = self.client.post(
            f"/api/datasets/{dataset_id}/clean",
            json={
                "action": "fillna",
                "params": {"method": "mean", "columns": ["销量", "销售额"]},
            },
        )
        self.assertTrue(response.get_json()["ok"], response.get_json())
        profiles = {item["name"]: item for item in self.summary(dataset_id)["profiles"]}
        self.assertEqual(profiles["销量"]["missing"], 0)
        self.assertEqual(profiles["销售额"]["missing"], 0)
        self.assertAlmostEqual(profiles["销量"]["mean"], 9.5, places=3)

        response = self.client.post(
            f"/api/datasets/{dataset_id}/clean",
            json={"action": "drop_columns", "params": {"columns": ["备注"]}},
        )
        self.assertTrue(response.get_json()["ok"])
        names = [item["name"] for item in self.summary(dataset_id)["columns"]]
        self.assertNotIn("备注", names)

        response = self.client.post(
            f"/api/datasets/{dataset_id}/clean",
            json={
                "action": "convert",
                "params": {"column": "编号", "target": "int", "on_error": "coerce"},
            },
        )
        self.assertTrue(response.get_json()["ok"], response.get_json())
        profiles = {item["name"]: item for item in self.summary(dataset_id)["profiles"]}
        self.assertEqual(profiles["编号"]["type"], "整数")

        response = self.client.post(
            f"/api/datasets/{dataset_id}/clean",
            json={"action": "drop_columns", "params": {"columns": ["编号"]}},
        )
        self.assertTrue(response.get_json()["ok"])

        response = self.client.post(f"/api/datasets/{dataset_id}/undo")
        self.assertTrue(response.get_json()["changed"])
        names = [item["name"] for item in self.summary(dataset_id)["columns"]]
        self.assertIn("编号", names)

        response = self.client.post(f"/api/datasets/{dataset_id}/reset")
        self.assertTrue(response.get_json()["changed"])
        summary = self.summary(dataset_id)
        self.assertEqual(summary["base_rows"], 6)
        self.assertEqual(summary["meta"]["column_count"], 8)
        self.assertFalse(summary["can_undo"])

    def test_04_filters_sorts_and_charts(self):
        dataset_id, _ = self.import_csv("图表测试")
        response = self.client.post(
            f"/api/datasets/{dataset_id}/filters",
            json={
                "filters": [
                    {"column": "地区", "op": "eq", "value": "华东"},
                ],
                "logic": "AND",
                "sorts": [{"column": "销售额", "order": "desc"}],
            },
        )
        self.assertTrue(response.get_json()["ok"], response.get_json())
        rows = self.client.get(f"/api/datasets/{dataset_id}/rows?page=1&page_size=50").get_json()
        self.assertEqual(rows["total"], 2)
        self.assertEqual(rows["rows"][0]["产品"], "显示器")

        chart_specs = [
            {"type": "bar", "x": "地区", "y": ["销售额"], "agg": "sum"},
            {"type": "line", "x": "日期", "y": ["销售额"], "agg": "sum"},
            {"type": "pie", "x": "产品", "y": ["销售额"], "agg": "sum"},
            {"type": "scatter", "x": "单价", "y": ["销售额"], "agg": "sum"},
        ]
        for spec in chart_specs:
            payload = dict(spec)
            payload.update({"sort": "auto", "top_n": 0})
            response = self.client.post(f"/api/datasets/{dataset_id}/chart", json=payload)
            self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
            body = response.get_json()
            self.assertTrue(body["ok"])
            self.assertIn("series", body["option"])

        response = self.client.post(
            f"/api/datasets/{dataset_id}/chart",
            json={"type": "bar", "x": "地区", "y": ["地区"], "agg": "sum"},
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("同一个字段", response.get_json()["error"])

        response = self.client.post(
            f"/api/datasets/{dataset_id}/chart",
            json={"type": "bar", "x": "地区", "y": ["产品"], "agg": "sum"},
        )
        self.assertEqual(response.status_code, 400)

    def test_05_exports(self):
        dataset_id, _ = self.import_csv("导出测试")
        response = self.client.get(f"/api/datasets/{dataset_id}/export.xlsx?scope=view")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data[:2], b"PK")
        workbook = load_workbook(io.BytesIO(response.data))
        self.assertIn("数据", workbook.sheetnames)
        self.assertIn("导出说明", workbook.sheetnames)
        self.assertEqual(workbook["数据"].max_row, 7)

        response = self.client.post(
            f"/api/datasets/{dataset_id}/chart/export",
            json={"type": "bar", "x": "地区", "y": ["销售额"], "agg": "sum"},
        )
        self.assertEqual(response.status_code, 200, response.data[:120])
        self.assertEqual(response.data[:8], b"\x89PNG\r\n\x1a\n")
        image = Image.open(io.BytesIO(response.data))
        self.assertGreater(image.width, 600)
        self.assertGreater(image.height, 400)

    def test_06_validation_and_pages(self):
        response = upload_file(self.client, b"abc", "bad.txt")
        self.assertEqual(response.status_code, 400)

        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertIn("从一份表格开始", response.get_data(as_text=True))

        response = self.client.get("/history")
        self.assertEqual(response.status_code, 200)
        self.assertIn("历史记录", response.get_data(as_text=True))

        response = self.client.get("/workspace/not-a-valid-id")
        self.assertEqual(response.status_code, 404)

    def test_99_history_and_delete(self):
        dataset_id, _ = self.import_csv("删除测试")
        history = self.client.get("/history").get_data(as_text=True)
        self.assertIn("上传", history)

        response = self.client.delete(f"/api/datasets/{dataset_id}")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.client.get(f"/workspace/{dataset_id}").status_code, 404)

        response = self.client.post("/history/clear", follow_redirects=True)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(application.database.get_stats()["total"], 0)


if __name__ == "__main__":
    unittest.main(verbosity=2)
