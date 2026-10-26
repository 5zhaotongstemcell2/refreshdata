"""生成软件著作权申请材料。

使用 python-docx 创建：
01 申请表信息稿
02 源程序（前 30 页 + 后 30 页）
03 软件说明书
04 申请材料清单与办理说明
另生成源程序完整文本，便于归档。
"""
from __future__ import annotations

import math
import os
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import copyright_text as text

BASE = Path(__file__).resolve().parent
OUT = BASE / "软件著作权申请材料"
FONT_CN = "Microsoft YaHei"
FONT_CODE = "Consolas"
SOFTWARE_HEADER = f"{text.SOFTWARE_NAME} {text.VERSION}"

SOURCE_ORDER = [
    "app.py",
    "config.py",
    "database.py",
    "deps.py",
    "services/importer.py",
    "services/storage.py",
    "services/serialize.py",
    "services/cleaner.py",
    "services/filters.py",
    "services/charts.py",
    "services/exporter.py",
    "services/history.py",
    "templates/base.html",
    "templates/index.html",
    "templates/workspace.html",
    "templates/history.html",
    "templates/error.html",
    "static/css/app.css",
    "static/js/common.js",
    "static/js/index.js",
    "static/js/workspace.js",
    "static/js/history.js",
]


# ---------------------------------------------------------------------------
# 通用格式工具
# ---------------------------------------------------------------------------


def set_run_font(
    run,
    name: str = FONT_CN,
    size: float = 10.5,
    bold: bool = False,
    color: str | None = None,
    italic: bool = False,
):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), FONT_CN if name != FONT_CODE else FONT_CN)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=60, start=80, bottom=60, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_field(paragraph, field: str, size: float = 9):
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), f" {field} ")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), FONT_CN)
    rfonts.set(qn("w:hAnsi"), FONT_CN)
    rfonts.set(qn("w:eastAsia"), FONT_CN)
    rpr.append(rfonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rpr.append(sz)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = "1"
    run.append(t)
    fld.append(run)
    paragraph._p.append(fld)


def set_update_fields(doc: Document):
    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def configure_section(section, top=2.2, bottom=2.0, left=2.4, right=2.4):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(top)
    section.bottom_margin = Cm(bottom)
    section.left_margin = Cm(left)
    section.right_margin = Cm(right)
    section.header_distance = Cm(0.9)
    section.footer_distance = Cm(0.8)


def clear_header_footer(section):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    for paragraph in section.header.paragraphs:
        for run in paragraph.runs:
            run.text = ""
    for paragraph in section.footer.paragraphs:
        for run in paragraph.runs:
            run.text = ""


def clear_header_style_tabs(doc: Document):
    """去掉 Header 样式的默认居中/右制表位，避免页眉右侧文字落在页面中间。"""
    try:
        doc.styles["Header"].paragraph_format.tab_stops.clear_all()
    except KeyError:
        pass


def header_left_right(
    section, left: str, right: str, size: float = 9, tab_cm: float = 16.2
):
    section.header.is_linked_to_previous = False
    paragraph = section.header.paragraphs[0]
    paragraph.text = ""
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.tab_stops.clear_all()
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Cm(tab_cm), WD_TAB_ALIGNMENT.RIGHT
    )
    run = paragraph.add_run(left)
    set_run_font(run, FONT_CN, size, color="4B5563")
    run = paragraph.add_run("\t" + right)
    set_run_font(run, FONT_CN, size, color="4B5563")


def footer_page_numbers(section, size: float = 9):
    section.footer.is_linked_to_previous = False
    paragraph = section.footer.paragraphs[0]
    paragraph.text = ""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("第 ")
    set_run_font(run, FONT_CN, size, color="4B5563")
    add_page_field(paragraph, "PAGE", size)
    run = paragraph.add_run(" 页 / 共 ")
    set_run_font(run, FONT_CN, size, color="4B5563")
    add_page_field(paragraph, "NUMPAGES", size)
    run = paragraph.add_run(" 页")
    set_run_font(run, FONT_CN, size, color="4B5563")


def add_table(doc, headers, rows, widths, font_size=9.5, header_fill="E8ECEF"):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        cell = header_cells[index]
        cell.width = Cm(widths[index])
        set_cell_margins(cell)
        set_cell_shading(cell, header_fill)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(str(header))
        set_run_font(run, FONT_CN, font_size, bold=True, color="1F2933")
    repeat_table_header(table.rows[0])
    for row_data in rows:
        row = table.add_row()
        prevent_row_split(row)
        for index, value in enumerate(row_data):
            cell = row.cells[index]
            cell.width = Cm(widths[index])
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.15
            run = paragraph.add_run(str(value))
            set_run_font(run, FONT_CN, font_size, color="1F2933")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_body(doc, content: str, indent=True):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.48
    if indent:
        paragraph.paragraph_format.first_line_indent = Pt(22)
    run = paragraph.add_run(content)
    set_run_font(run, FONT_CN, 11, color="1F2933")
    return paragraph


def add_bullet(doc, content: str, index: int | None = None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.45
    paragraph.paragraph_format.left_indent = Pt(12)
    paragraph.paragraph_format.first_line_indent = Pt(-12)
    prefix = f"（{index}）" if index is not None else "· "
    run = paragraph.add_run(prefix + content)
    set_run_font(run, FONT_CN, 10.5, color="1F2933")
    return paragraph


def add_steps(doc, steps):
    for index, step in enumerate(steps, start=1):
        add_bullet(doc, step, index)


def add_code(doc, lines):
    for line in lines:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.05
        paragraph.paragraph_format.left_indent = Pt(10)
        ppr = paragraph._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "F4F5F6")
        ppr.append(shd)
        run = paragraph.add_run(line)
        set_run_font(run, FONT_CODE, 8.8, color="1F2933")


def add_heading(doc, content: str, level: int):
    paragraph = doc.add_heading(content, level=level)
    paragraph.paragraph_format.keep_with_next = True
    if level == 1:
        paragraph.paragraph_format.space_before = Pt(14)
        paragraph.paragraph_format.space_after = Pt(5)
        for run in paragraph.runs:
            set_run_font(run, FONT_CN, 15, bold=True, color="1F3B4D")
    else:
        paragraph.paragraph_format.space_before = Pt(9)
        paragraph.paragraph_format.space_after = Pt(3)
        for run in paragraph.runs:
            set_run_font(run, FONT_CN, 12.5, bold=True, color="2F4F63")
    return paragraph


# ---------------------------------------------------------------------------
# 源程序
# ---------------------------------------------------------------------------


def visual_width(value: str) -> int:
    return sum(2 if ord(char) > 127 else 1 for char in value)


def wrap_code_line(line: str, limit: int = 104) -> list[str]:
    line = line.replace("\t", "    ").rstrip()
    if not line:
        return []
    chunks = []
    current = ""
    width = 0
    for char in line:
        char_width = 2 if ord(char) > 127 else 1
        if width + char_width > limit and current:
            chunks.append(current)
            current = "    " + char
            width = 4 + char_width
        else:
            current += char
            width += char_width
    if current:
        chunks.append(current)
    return chunks


def marker_for(relative_path: str) -> str:
    suffix = Path(relative_path).suffix.lower()
    if suffix == ".py":
        return f"# ==================== FILE: {relative_path} ===================="
    if suffix in (".html", ".htm"):
        return f"<!-- ==================== FILE: {relative_path} ==================== -->"
    return f"/* ==================== FILE: {relative_path} ==================== */"


def collect_source():
    original_lines = 0
    formatted = []
    for relative in SOURCE_ORDER:
        path = BASE / relative
        content = path.read_text(encoding="utf-8-sig")
        lines = content.splitlines()
        original_lines += len(lines)
        formatted.append(marker_for(relative))
        for line in lines:
            formatted.extend(wrap_code_line(line))
    return original_lines, formatted


def add_source_lines(doc, lines, start_page: int, page_count: int):
    body_area = doc.sections[0]
    for page_index in range(page_count):
        chunk = lines[(start_page + page_index) * 50 : (start_page + page_index + 1) * 50]
        if len(chunk) < 50:
            raise ValueError("源程序行数不足，无法组成完整 50 行页面。")
        for line_index, line in enumerate(chunk):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = Pt(13.2)
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.left_indent = Pt(0)
            paragraph.paragraph_format.widow_control = False
            if line_index == 0 and page_index > 0:
                paragraph.paragraph_format.page_break_before = True
            run = paragraph.add_run(line if line else " ")
            set_run_font(run, FONT_CODE, 8, color="111827")


def build_source_docx(lines, path: Path):
    doc = Document()
    clear_header_style_tabs(doc)
    section = doc.sections[0]
    configure_section(section, top=1.7, bottom=1.7, left=1.6, right=1.6)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)
    style = doc.styles["Normal"]
    style.font.name = FONT_CODE
    style.font.size = Pt(8)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    header_left_right(
        section, SOFTWARE_HEADER, "源程序（前30页+后30页）", size=9, tab_cm=17.8
    )
    footer_page_numbers(section, size=9)
    add_source_lines(doc, lines, 0, 60)
    set_update_fields(doc)
    doc.core_properties.title = f"{text.SOFTWARE_NAME} {text.VERSION} 源程序"
    doc.core_properties.author = text.SOFTWARE_SHORT
    doc.core_properties.subject = "计算机软件著作权登记源程序鉴别材料"
    doc.save(path)


def build_source_txt(path: Path):
    parts = []
    for relative in SOURCE_ORDER:
        content = (BASE / relative).read_text(encoding="utf-8-sig")
        parts.append(marker_for(relative))
        parts.append(content.rstrip("\n"))
        parts.append("")
    path.write_text("\n".join(parts) + "\n", encoding="utf-8")


# ---------------------------------------------------------------------------
# 说明书
# ---------------------------------------------------------------------------


def build_manual(path: Path):
    doc = Document()
    clear_header_style_tabs(doc)
    section = doc.sections[0]
    configure_section(section, top=2.2, bottom=2.0, left=2.4, right=2.4)
    header_left_right(section, SOFTWARE_HEADER, "软件说明书", size=9)
    footer_page_numbers(section, size=9)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_CN
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal.paragraph_format.line_spacing = 1.48
    normal.paragraph_format.space_after = Pt(3)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(12)
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run(text.SOFTWARE_NAME)
    set_run_font(run, FONT_CN, 22, bold=True, color="1F3B4D")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(10)
    run = subtitle.add_run(f"{text.VERSION}  软件说明书")
    set_run_font(run, FONT_CN, 15, bold=False, color="2F4F63")

    info_rows = [
        ["软件名称", f"{text.SOFTWARE_NAME} {text.VERSION}"],
        ["软件简称", text.SOFTWARE_SHORT],
        ["著作权人", "【著作权人名称】"],
        ["文档版本", "1.0"],
        ["编制日期", text.DOC_DATE],
    ]
    add_table(doc, ["项目", "内容"], info_rows, [4.0, 11.8], font_size=9.5)

    toc_heading = doc.add_paragraph()
    toc_heading.paragraph_format.space_before = Pt(4)
    toc_heading.paragraph_format.space_after = Pt(2)
    run = toc_heading.add_run("目录")
    set_run_font(run, FONT_CN, 13, bold=True, color="1F3B4D")

    for entry in text.TOC:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        run = paragraph.add_run(entry)
        set_run_font(run, FONT_CN, 9.5, color="1F2933")

    first_section = True

    for kind, payload in text.MANUAL:
        if kind == "h1":
            heading = add_heading(doc, payload, 1)
            if first_section:
                heading.paragraph_format.page_break_before = True
                first_section = False
        elif kind == "h2":
            add_heading(doc, payload, 2)
        elif kind == "p":
            add_body(doc, payload)
        elif kind == "bullets":
            for index, item in enumerate(payload, start=1):
                add_bullet(doc, item, index)
        elif kind == "steps":
            add_steps(doc, payload)
        elif kind == "table":
            add_table(doc, payload["headers"], payload["rows"], payload["widths"])
        elif kind == "code":
            add_code(doc, payload)
        elif kind == "faq":
            for question, answer in payload:
                q = doc.add_paragraph()
                q.paragraph_format.space_before = Pt(4)
                q.paragraph_format.space_after = Pt(1)
                q.paragraph_format.keep_with_next = True
                run = q.add_run("问：" + question)
                set_run_font(run, FONT_CN, 10.8, bold=True, color="1F3B4D")
                a = doc.add_paragraph()
                a.paragraph_format.space_before = Pt(0)
                a.paragraph_format.space_after = Pt(4)
                a.paragraph_format.line_spacing = 1.4
                a.paragraph_format.left_indent = Pt(12)
                run = a.add_run("答：" + answer)
                set_run_font(run, FONT_CN, 10.5, color="1F2933")
        else:
            raise ValueError(f"未知的说明书块：{kind}")

    # 最后一个表格后会留下一个空段落，若页面刚好排满会产生空白页
    body = doc.element.body
    for paragraph in reversed(list(body.findall(qn("w:p")))):
        if "".join(paragraph.itertext()).strip():
            break
        body.remove(paragraph)

    set_update_fields(doc)
    doc.core_properties.title = f"{text.SOFTWARE_NAME} {text.VERSION} 软件说明书"
    doc.core_properties.author = text.SOFTWARE_SHORT
    doc.core_properties.subject = "计算机软件著作权登记文档鉴别材料"
    doc.save(path)


# ---------------------------------------------------------------------------
# 申请表
# ---------------------------------------------------------------------------


def form_table(doc, rows, widths=(4.2, 12.4)):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for label, value in rows:
        row = table.add_row()
        prevent_row_split(row)
        left, right = row.cells
        left.width = Cm(widths[0])
        right.width = Cm(widths[1])
        set_cell_margins(left, 80, 100, 80, 100)
        set_cell_margins(right, 80, 100, 80, 100)
        set_cell_shading(left, "F3F4F6")
        p = left.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(label)
        set_run_font(run, FONT_CN, 9.5, bold=True, color="1F2933")
        p = right.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.25
        run = p.add_run(str(value))
        set_run_font(run, FONT_CN, 9.5, color="1F2933")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build_form(path: Path):
    doc = Document()
    section = doc.sections[0]
    configure_section(section, top=2.0, bottom=2.0, left=2.2, right=2.2)
    footer_page_numbers(section, size=9)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(4)
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("计算机软件著作权登记申请表")
    set_run_font(run, FONT_CN, 19, bold=True, color="1F3B4D")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(8)
    run = subtitle.add_run("（信息填写稿 · 用于线上系统填报）")
    set_run_font(run, FONT_CN, 11, color="6B7280")

    note = doc.add_paragraph()
    note.paragraph_format.space_after = Pt(6)
    run = note.add_run(
        "说明：本文件是线上填报前的信息底稿，不是中国版权保护中心的正式申请确认签章页。"
        "请在“中国版权保护中心”软件登记系统在线填写后，打印系统生成的申请确认签章页，"
        "由申请人签字或加盖公章，再按要求扫描上传。"
    )
    set_run_font(run, FONT_CN, 9.5, color="6B7280")

    add_heading(doc, "一、软件基本信息", 2)
    form_table(doc, [
        ("软件全称", text.FORM["软件全称"]),
        ("软件简称", text.FORM["软件简称"]),
        ("版本号", text.FORM["版本号"]),
        ("软件作品说明", text.FORM["软件作品说明"]),
        ("开发完成日期", text.FORM["开发完成日期"]),
        ("发表状态", text.FORM["发表状态"]),
        ("开发方式", text.FORM["开发方式"]),
        ("软件分类", text.FORM["软件分类"]),
    ])

    add_heading(doc, "二、著作权人信息", 2)
    applicant_keys = [
        "著作权人名称",
        "主体类别",
        "证件类型",
        "证件号码",
        "国籍",
        "省份/城市",
        "通讯地址",
        "邮政编码",
        "联系人",
        "联系电话",
        "电子邮箱",
    ]
    form_table(doc, [(key, text.APPLICANT[key]) for key in applicant_keys])

    add_heading(doc, "三、权利说明", 2)
    form_table(doc, [
        ("权利取得方式", text.FORM["权利取得方式"]),
        ("权利范围", text.FORM["权利范围"]),
        ("原登记号", "无（首次申请）"),
        ("变更/补充证明号", "无"),
    ])

    add_heading(doc, "四、软件鉴别材料", 2)
    form_table(doc, [
        ("交存方式", "一般交存"),
        ("源程序", "前、后各连续 30 页，共 60 页；每页 50 行；页眉标注软件名称及版本号"),
        ("文档", "《数据洞察可视化分析软件 V1.0 软件说明书》全部页面；每页不少于 30 行"),
        ("源程序量", text.FORM["源程序量"]),
        ("页眉版本一致性", f"源程序与说明书页眉均为“{SOFTWARE_HEADER}”，与申请表版本号一致"),
    ])

    add_heading(doc, "五、软件功能和技术特点", 2)
    form_table(doc, [
        ("硬件环境", text.FORM["硬件环境"]),
        ("软件环境", text.FORM["软件环境"]),
        ("编程语言", text.FORM["编程语言"]),
        ("源程序量", text.FORM["源程序量"]),
        ("开发目的", text.FORM["开发目的"]),
        ("面向行业", text.FORM["面向行业"]),
        ("主要功能", text.FORM["主要功能"]),
        ("技术特点", text.FORM["技术特点"]),
    ], widths=(3.8, 12.8))

    add_heading(doc, "六、申请办理与签章", 2)
    form_table(doc, [
        ("申请办理方式", text.APPLICANT["办理方式"]),
        ("代理机构", text.APPLICANT["代理机构"]),
        ("申请人承诺", "申请人承诺所提交的软件为独立开发（或已取得合法授权），提交材料真实、合法、有效，愿承担相应法律责任。"),
        ("申请人签章", "（自然人签名 / 单位公章）\n\n年    月    日"),
    ])

    add_heading(doc, "七、随附材料清单", 2)
    form_table(doc, [
        ("申请表", "线上填写并打印申请确认签章页，签字/盖章后扫描为 PDF"),
        ("源程序", "《源程序（前30页+后30页）》PDF，共 60 页"),
        ("文档", "《软件说明书》PDF，全部页面"),
        ("身份证明", "自然人身份证 / 法人营业执照或统一社会信用代码证书扫描件"),
        ("权利归属证明", "独立开发无需提交；委托、合作、任务开发需提交相应合同或任务书"),
    ])

    doc.add_paragraph()
    run = doc.add_paragraph().add_run(
        "注：表中【】内容为待填写项；“发表状态”如已发表，请填写首次发表日期和地点，并删除“未发表”表述。"
        "软件名称、版本号、页眉、说明书和源程序必须保持一致。"
    )
    set_run_font(run, FONT_CN, 9, color="6B7280")

    set_update_fields(doc)
    doc.core_properties.title = f"{text.SOFTWARE_NAME} {text.VERSION} 软件著作权登记申请表（信息稿）"
    doc.save(path)


# ---------------------------------------------------------------------------
# 材料清单与办理说明
# ---------------------------------------------------------------------------


def build_checklist(path: Path):
    doc = Document()
    section = doc.sections[0]
    configure_section(section, top=2.2, bottom=2.0, left=2.4, right=2.4)
    footer_page_numbers(section, size=9)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("软件著作权登记材料清单与办理说明")
    set_run_font(run, FONT_CN, 19, bold=True, color="1F3B4D")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(8)
    run = subtitle.add_run(f"{text.SOFTWARE_NAME} {text.VERSION}")
    set_run_font(run, FONT_CN, 11, color="6B7280")

    add_heading(doc, "一、本套文件包含的内容", 1)
    add_table(doc, ["序号", "文件", "用途"], [
        ["01", "01_软件著作权登记申请表（信息稿）.docx / .pdf", "线上填报的信息底稿，不作为正式签章页"],
        ["02", "02_源程序（前30页+后30页）.docx / .pdf", "源程序鉴别材料，提交 PDF"],
        ["03", "03_软件说明书.docx / .pdf", "文档鉴别材料，提交 PDF"],
        ["04", "04_申请材料清单与办理说明.docx / .pdf", "本文件，办理指引"],
        ["05", "05_源程序_完整版.txt", "完整源程序归档，不作为提交页数依据"],
    ], [1.4, 8.6, 6.4], font_size=9.5)

    add_heading(doc, "二、还需要申请人自行提供的材料", 1)
    add_table(doc, ["材料", "说明"], [
        ["身份证明", "自然人：身份证正反面照片/扫描件；法人：加载统一社会信用代码的营业执照副本或证书原件的照片/扫描件。"],
        ["申请确认签章页", "在版权中心线上系统填写后打印，自然人本人签名或单位加盖鲜章，再扫描为 PDF 上传。"],
        ["权利归属证明", "独立开发、原始取得的一般不需要。委托开发、合作开发、下达任务开发、受让、继承等情形，按官方要求提交合同、任务书或其他证明。"],
        ["代理授权文件", "如委托代理机构办理，按代理机构要求提供授权委托书。"],
    ], [4.2, 12.2], font_size=9.5)

    add_heading(doc, "三、格式与一致性检查", 1)
    add_p = lambda content: add_body(doc, content)
    add_p("（1）申请文件应纵向排版，文字从左向右排列。源程序和说明书均使用 A4 纸张。")
    add_p("（2）源程序提交前、后各连续 30 页，共 60 页；每页 50 行，页眉标注软件名称和版本号，页脚标注页码。本套源程序第 1–30 页为前 30 页，第 31–60 页为后 30 页。")
    add_p("（3）说明书每页不少于 30 行；本套说明书按连续正文排版，不含大面积空白页。如自行替换内容，请保持每页文字密度。")
    add_p(f"（4）软件全称统一为“{text.SOFTWARE_NAME}”，版本号统一为“{text.VERSION}”。申请表、源程序页眉、说明书页眉、封面和导出文件名必须完全一致，包括 V 的有无。")
    add_p("（5）源程序页眉不得出现与申请表不一致的简称；如改软件名称，请同时修改申请表、源程序页眉、说明书页眉和文档内所有引用。")
    add_p("（6）申请表在线填写后，申请确认签章页不得擅自更改内容、格式和打印比例；签字/盖章须为鲜章或本人签名。")

    add_heading(doc, "四、线上办理步骤", 1)
    add_steps(doc, [
        "访问中国版权保护中心软件登记系统（register.ccopyright.com.cn），注册并完成身份认证。",
        "选择“软件著作权登记”，按系统提示填写软件基本信息、著作权人信息、权利说明和技术信息。",
        "填写完成后在线打印申请确认签章页，由申请人签字或加盖单位公章，扫描为 PDF。",
        "上传申请确认签章页、源程序 PDF、说明书 PDF 和身份证明文件。",
        "确认材料齐全后提交申请，关注系统受理和补正通知。",
        "登记完成后，在用户中心下载软件著作权登记电子证书。",
    ])
    add_p("具体界面、字段和上传要求以系统当前版本为准。官方指南页面：https://www.ccopyright.com/index.php?optionid=1080 ；填表说明：https://www.ccopyright.com/index.php?optionid=1081 。")
    add_p("如系统要求提交“前、后各连续 30 页”的源程序和文档，可直接使用本套 60 页源程序 PDF；说明书不足 60 页时按要求提交全部页面。")

    add_heading(doc, "五、提交前逐项核对", 1)
    add_table(doc, ["检查项", "确认"], [
        ["软件全称、简称、版本号已填写，且与所有页眉一致", "□"],
        ["开发完成日期不晚于申请日期", "□"],
        ["发表状态选择正确；已发表则填写首次发表日期和地点", "□"],
        ["开发方式选择正确；合作/委托/任务开发已准备权利归属证明", "□"],
        ["著作权人名称、证件类型、证件号码与身份证明一致", "□"],
        ["源程序 60 页、每页 50 行、页码连续、页眉有软件名称和版本", "□"],
        ["说明书页面完整、每页文字不少于 30 行、页眉版本一致", "□"],
        ["申请确认签章页已签字/盖章并扫描为清晰 PDF", "□"],
        ["文件格式、大小符合系统上传要求", "□"],
        ["已保存一份完整电子材料备份", "□"],
    ], [14.0, 2.4], font_size=9.5)

    add_heading(doc, "六、说明", 1)
    add_body(doc, "本套文件根据中国版权保护中心公开的《计算机软件著作权登记指南》整理，用于协助准备申请材料，不构成法律意见，也不替代受理机构的审查。具体要求、收费（如有）、办理时限和系统字段以中国版权保护中心及其登记系统的最新规定为准。")
    add_body(doc, "如申请人是单位，且软件由员工在职务范围内开发，建议同时保留劳动合同、项目立项或职务开发说明等内部资料，以备审查或后续维权使用；是否需要在本次登记中提交，以受理机构要求为准。")
    add_body(doc, "本套材料中的【著作权人名称】【证件号码】【地址】【联系方式】等占位内容，请在提交前替换为真实信息，并删除所有【】标记和“信息稿”字样。")

    set_update_fields(doc)
    doc.core_properties.title = "软件著作权登记材料清单与办理说明"
    doc.save(path)


# ---------------------------------------------------------------------------
# 主流程
# ---------------------------------------------------------------------------


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    original_lines, formatted_lines = collect_source()
    print(f"源文件总行数：{original_lines}")
    print(f"格式化后可提交行数：{len(formatted_lines)}")
    if len(formatted_lines) < 3000:
        raise SystemExit("源程序不足 3000 行，无法组成 60 页提交版。")

    excerpt = formatted_lines[:1500] + formatted_lines[-1500:]
    if len(excerpt) != 3000:
        raise SystemExit("源程序摘录行数不是 3000。")

    build_source_docx(excerpt, OUT / "02_源程序（前30页+后30页）.docx")
    build_source_txt(OUT / "05_源程序_完整版.txt")
    build_manual(OUT / "03_软件说明书.docx")
    build_form(OUT / "01_软件著作权登记申请表（信息稿）.docx")
    build_checklist(OUT / "04_申请材料清单与办理说明.docx")

    print("已生成：")
    for item in sorted(OUT.iterdir()):
        print(f"  {item.name}  {item.stat().st_size:,} bytes")


if __name__ == "__main__":
    main()
